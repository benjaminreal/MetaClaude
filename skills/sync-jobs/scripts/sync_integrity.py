"""Acquisition and two-surface consistency checks owned by sync-jobs."""
import json
import os
import re
from pathlib import Path
import openpyxl

ID = re.compile(r'/jobs/view/(\d+)')

def saved_records(paths):
    records=[]
    for path in paths or []:
        data=json.loads(Path(path).read_text())
        if not isinstance(data,dict) or not isinstance(data.get('jobs'),list):
            raise ValueError(f'{path}: saved index requires a metadata envelope')
        if data.get('error') not in (None,''):
            raise ValueError(f'{path}: acquisition failed: {data["error"]}')
        jobs=data['jobs']; ids=[str(j.get('id','')) for j in jobs]
        if any(not i.isdigit() for i in ids) or len(set(ids))!=len(ids):
            raise ValueError(f'{path}: missing, invalid or duplicate index IDs')
        count=data.get('count');total=data.get('total')
        if type(count)!=int or type(total)!=int or count!=len(jobs) or total<count:
            raise ValueError(f'{path}: invalid counts')
        if total!=count:
            raise ValueError(f'{path}: incomplete index (reported={total}, retrieved={count})')
        if data.get('pagination_complete') is False or data.get('termination') in {'error','guard_limit','unknown'}:
            raise ValueError(f'{path}: pagination did not complete')
        records.extend(jobs)
    return records

def surfaces(tracker,root,postings):
    root=Path(root).resolve(); rows={};archives={};tracker_ids=set()
    wb=openpyxl.load_workbook(tracker,read_only=True,data_only=True)
    ws=wb['Jobs'];headers={str(c.value):i for i,c in enumerate(ws[1]) if c.value is not None}
    for values in ws.iter_rows(min_row=2,values_only=True):
        row={h:values[i] if i<len(values) else None for h,i in headers.items()}
        tid=str(row.get('Tracker ID') or '')
        if tid and tid in tracker_ids:raise ValueError(f'Duplicate Tracker ID {tid}; resolve before sync')
        if tid:tracker_ids.add(tid)
        match=ID.search(str(row.get('Link puesto linkedin') or ''))
        if not match:continue
        jid=match[1]
        if not tid:raise ValueError(f'LinkedIn ID {jid} has no Tracker ID; manual reconciliation required')
        rows.setdefault(jid,[]).append(row)
    wb.close()
    for jid, group in list(rows.items()):
        if len(group)==1:
            rows[jid]=group[0]
            continue
        canonical=[r for r in group if str(r.get('Estatus') or '').strip().lower()!='duplicate']
        if len(canonical)!=1:raise ValueError(f'Duplicate tracker LinkedIn ID {jid}; resolve before sync')
        keeper=str(canonical[0]['Tracker ID'])
        for row in group:
            if row is canonical[0]:continue
            keepers=set(re.findall(r'duplicate\s+of\s+(J-\d+)',str(row.get('Comentarios') or ''),re.I))
            if keepers!={keeper}:raise ValueError(f'Duplicate row {row["Tracker ID"]} lacks matching keeper {keeper} for LinkedIn ID {jid}')
        rows[jid]=canonical[0]
    candidates=set(Path(postings).glob('*.md'))
    for row in rows.values():
        if row.get('JD File'):
            p=(root/str(row['JD File'])).resolve()
            if not p.is_relative_to(root):raise ValueError('JD File escapes selected workspace')
            if p.is_file():candidates.add(p)
    for p in candidates:
        for line in p.read_text().splitlines():
            if line.startswith('Source:'):
                match=ID.search(line)
                if match:
                    jid=match[1]
                    if jid in archives:
                        prior=archives[jid]
                        try:same=os.path.samefile(prior,p)
                        except OSError:same=prior.resolve()==p.resolve()
                        if not same:raise ValueError(f'Ambiguous archives for LinkedIn ID {jid}')
                    archives[jid]=p
                break
    # A tracker row may deliberately bind a LinkedIn Saved identity to a JD
    # captured from the employer's own page.  Preserve that higher-authority
    # source instead of treating it as a missing archive, but only through the
    # row's exact JD File path.  One physical archive cannot satisfy two IDs.
    bound_paths={str(path.resolve()):jid for jid,path in archives.items()}
    for jid,row in rows.items():
        if jid in archives or not row.get('JD File'):
            continue
        path=(root/str(row['JD File'])).resolve()
        if not path.is_file():
            continue
        prior=bound_paths.get(str(path))
        if prior and prior!=jid:
            # A historical repost may point at another LinkedIn identity's
            # archive.  Do not manufacture a second canonical binding; leave
            # this ID visible as recovery instead.
            continue
        source=next((line for line in path.read_text().splitlines() if line.startswith('Source:')),None)
        if source and not ID.search(source):
            archives[jid]=path
            bound_paths[str(path)]=jid
    return rows,archives

def archived_record(path,jid,status='Saved'):
    text=Path(path).read_text();parts=text.split('## About the Job',1)
    if len(parts)!=2 or not parts[1].strip():raise ValueError(f'Archive {jid} lacks a recoverable full description; manual reconciliation required')
    lines=parts[0].splitlines();fields={}
    for line in lines:
        if ': ' in line:
            key,value=line.split(': ',1);fields[key]=value
    if not lines or not lines[0].startswith('# ') or not fields.get('Company'):
        raise ValueError(f'Archive {jid} lacks identity metadata; manual reconciliation required')
    return {'id':jid,'title':lines[0][2:],'company':fields['Company'],'location':fields.get('Location','Unknown'),'type':fields.get('Type','Unknown'),'posted':fields.get('Posted',''),'description':parts[1].strip(),'url':f'https://www.linkedin.com/jobs/view/{jid}/','status':status}
