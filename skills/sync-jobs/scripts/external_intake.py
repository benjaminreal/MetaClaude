#!/usr/bin/env python3
"""Deterministic intake for owner-selected non-LinkedIn posting URLs."""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import tempfile
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

import openpyxl
from openpyxl.utils.cell import range_boundaries

import posting_compensation
import posting_eligibility
from acquisition_quality import validate_quality
from source_acquisition import source_record_hash, validate_bundle as validate_source_bundle
from sync_ingest import country_from_location, header_index, posted_yyyymm, sanitize

TRACKING_KEYS={"from","refid","source","eid","locale"}

def canonical_url(value: str) -> str:
    parsed=urlsplit(str(value).strip())
    if parsed.scheme.lower() not in {"http","https"} or not parsed.netloc:
        raise ValueError("direct posting URL must be absolute HTTP(S)")
    query=[]
    for key,val in parse_qsl(parsed.query,keep_blank_values=True):
        if key.lower().startswith("utm_") or key.lower() in TRACKING_KEYS: continue
        query.append((key,val))
    path=parsed.path or "/"
    return urlunsplit((parsed.scheme.lower(),parsed.netloc.lower(),path,urlencode(query),""))

def source_id(url: str) -> str:
    return "ext-"+hashlib.sha256(canonical_url(url).encode()).hexdigest()[:20]

def canonical_json(value) -> bytes:
    return json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(",",":")).encode()

def load_records(path: Path) -> tuple[list[dict], str | None]:
    payload=json.loads(path.read_text())
    if not isinstance(payload,dict) or not isinstance(payload.get("records"),list):
        raise ValueError("external input must contain records[]")
    bundle_sha=None
    if payload.get("schema") == "SourcePostingAcquisitionV2":
        validate_source_bundle(payload)
        if payload.get("complete") is not True or payload.get("quality_complete") is not True:
            raise ValueError("source acquisition bundle is incomplete; ingestion requires the original complete batch")
        bundle_sha=payload["bundle_sha256"]
    elif payload.get("schema") is not None:
        raise ValueError("unsupported external input schema")
    if payload.get("failures") not in (None,[]):
        raise ValueError("external input retains failures; ingestion requires a complete batch")
    records=[];seen=set()
    for raw in payload["records"]:
        if not isinstance(raw,dict): raise ValueError("external records must be objects")
        rec=dict(raw);url=rec.get("url")
        sid=source_id(url)
        if rec.get("source_id") not in (None,sid): raise ValueError(f"{sid}: supplied source_id does not match canonical URL")
        rec["source_id"]=sid;rec["id"]=sid
        for key in ("title","company","location","description"):
            if not isinstance(rec.get(key),str) or not rec[key].strip(): raise ValueError(f"{sid}: {key} is required")
        if sid in seen: raise ValueError(f"duplicate external source identity {sid}")
        seen.add(sid);quality=validate_quality(rec,require_pass=True)
        description_sha=hashlib.sha256(rec["description"].encode()).hexdigest()
        if rec.get("description_sha256") not in (None,description_sha):
            raise ValueError(f"{sid}: supplied description_sha256 does not match description bytes")
        supplied_source_hash=rec.get("source_record_sha256")
        expected_source_hash=source_record_hash(rec)
        if supplied_source_hash not in (None,expected_source_hash):
            raise ValueError(f"{sid}: supplied source_record_sha256 does not match source fields")
        if bundle_sha is None and rec.get("quality") not in (None,quality):
            raise ValueError(f"{sid}: supplied quality result is stale or fabricated")
        rec["description_sha256"]=description_sha
        rec["quality"]=quality
        rec["source_record_sha256"]=source_record_hash(rec)
        records.append(rec)
    if not records: raise ValueError("external input has no records")
    return records,bundle_sha

def source_sidecar(rec: dict, tracker_id: str, archive_path: str, bundle_sha: str | None) -> dict:
    """Retain source-neutral identity without overloading workbook columns."""
    return {
        "schema":"PostingSourceProvenanceV2",
        "version":"2.0",
        "tracker_id":tracker_id,
        "archive_path":archive_path,
        "source_id":rec["source_id"],
        "source_record_sha256":rec["source_record_sha256"],
        "acquisition_bundle_sha256":bundle_sha,
        "discovery_url":rec.get("supplied_url") or rec["url"],
        "final_url":rec.get("final_url") or rec["url"],
        "canonical_url":rec.get("canonical_url") or canonical_url(rec["url"]),
        "provider":rec.get("provider"),
        "ats":rec.get("ats"),
        "provider_job_id":rec.get("provider_job_id"),
        "requisition_id":rec.get("requisition_id"),
        "source_identity":rec.get("source_identity") or {"kind":"canonical_url","canonical_url":canonical_url(rec["url"])},
        "authentication_state":rec.get("authentication_state") or "unknown",
        "provenance":rec["provenance"],
        "attempts":rec.get("attempts",[]),
    }

def render_md(rec: dict) -> str:
    return (f"# {rec['title'].strip()}\n\nCompany: {rec['company'].strip()}\n"
            f"Location: {rec['location'].strip()}\n"
            f"Type: {str(rec.get('type') or rec.get('employmentType') or 'Unknown').strip()}\n"
            f"Posted: {posted_yyyymm(rec)}\nSource: {rec['url'].strip()}\n"
            f"Source ID: {rec['source_id']}\n\n## About the Job\n\n{rec['description'].strip()}\n")

def tracker_urls(ws,hdr):
    out={}
    cols=[hdr.get("Link puesto linkedin"),hdr.get("Link puesto empresa")]
    for row in range(2,ws.max_row+1):
        tid=str(ws.cell(row,hdr["Tracker ID"]).value or "")
        for col in cols:
            value=ws.cell(row,col).value if col else None
            if not value: continue
            try:key=canonical_url(str(value))
            except ValueError: continue
            item=(tid,row)
            if item not in out.setdefault(key,[]):out[key].append(item)
    return out

def archive_urls(postings: Path):
    out={}
    for path in postings.glob("*.md"):
        source=next((line.split(":",1)[1].strip() for line in path.read_text(errors="ignore").splitlines() if line.startswith("Source:")),None)
        if not source: continue
        try:key=canonical_url(source)
        except ValueError: continue
        out.setdefault(key,[]).append(path)
    return out

def max_tracker_id(ws,col):
    values=[]
    for row in range(2,ws.max_row+1):
        match=re.fullmatch(r"J-(\d+)",str(ws.cell(row,col).value or ""))
        if match: values.append(int(match.group(1)))
    return max(values,default=0)

def extend_table(ws):
    table=ws.tables.get("JobsTable")
    if not table:return None
    min_col,min_row,max_col,_=range_boundaries(table.ref)
    table.ref=f"{openpyxl.utils.get_column_letter(min_col)}{min_row}:{openpyxl.utils.get_column_letter(max_col)}{ws.max_row}"
    if table.autoFilter is not None:table.autoFilter.ref=table.ref
    return table.ref

def main(argv=None):
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--project-root",required=True);ap.add_argument("--tracker",required=True)
    ap.add_argument("--records-json",required=True);ap.add_argument("--dry-run",action="store_true")
    ap.add_argument("--backup-dir",default="/private/tmp")
    args=ap.parse_args(argv)
    root=Path(args.project_root).resolve();tracker=Path(args.tracker).resolve();postings=root/"JobPostings/postings";meta=root/"JobPostings/_meta"
    if not tracker.is_relative_to(root):ap.error("tracker must be inside project root")
    records,bundle_sha=load_records(Path(args.records_json).resolve());initial=hashlib.sha256(tracker.read_bytes()).hexdigest()
    wb_ro=openpyxl.load_workbook(tracker,read_only=True,data_only=True);ws_ro=wb_ro["Jobs"];hdr_ro={str(c.value).strip():i for i,c in enumerate(ws_ro[1],1) if c.value is not None}
    required=["Tracker ID","Previous Row","Puesto","Empresa","Pais","Estatus","Next Action","JD File","Link puesto linkedin","Link puesto empresa","Comentarios"]
    missing=[x for x in required if x not in hdr_ro]
    if missing:raise ValueError(f"tracker missing headers: {missing}")
    urls=tracker_urls(ws_ro,hdr_ro);wb_ro.close();archives=archive_urls(postings)
    pending=[];already=[]
    for rec in records:
        key=canonical_url(rec["url"]);t=urls.get(key,[]);a=archives.get(key,[])
        if len(t)>1 or len(a)>1:raise ValueError(f"{rec['source_id']}: ambiguous existing direct-source identity")
        if bool(t)!=bool(a):raise ValueError(f"{rec['source_id']}: partial tracker/archive state requires explicit recovery")
        (already if t else pending).append(rec)
    if not pending:
        print(json.dumps({"status":"ALREADY_PRESENT","written_count":0,"already_present_count":len(already),"source_ids":[r["source_id"] for r in already]},ensure_ascii=False,indent=2))
        return 0
    stamp=dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    stage=Path(args.backup_dir)/f"sync-jobs-external-{stamp}"
    stage.mkdir(parents=True,exist_ok=False)
    target=stage/tracker.name
    shutil.copy2(tracker,target)
    backup=stage/f"{tracker.stem}.backup{tracker.suffix}";shutil.copy2(tracker,backup)
    stage_postings=stage/"postings"
    stage_meta=stage/"meta"
    stage_postings.mkdir(parents=True,exist_ok=True)
    wb=openpyxl.load_workbook(target);ws=wb["Jobs"];hdr=header_index(ws);base=max_tracker_id(ws,hdr["Tracker ID"]);written=[]
    for offset,rec in enumerate(pending,1):
        tid=f"J-{base+offset:06d}";month=posted_yyyymm(rec)
        fname=f"{sanitize(rec['company'])}_{sanitize(rec['title'])}_{month}_{rec['source_id']}.md";path=stage_postings/fname
        archive_rel=f"JobPostings/postings/{fname}"
        if path.exists():raise ValueError(f"archive filename collision: {fname}")
        path.write_text(render_md(rec),encoding="utf-8")
        capture=rec["provenance"];binding={"captured_at":capture["captured_at"],"record_sha256":rec["source_record_sha256"],"source_id":rec["source_id"],"source_url":rec["url"],"description_sha256":rec["description_sha256"],"kind":"validated_external_acquisition_record"}
        source_meta=source_sidecar(rec,tid,archive_rel,bundle_sha)
        source_meta_path=stage_meta/"sources"/f"{tid}_PostingSourceProvenanceV2.json"
        posting_compensation.atomic_write(source_meta_path,json.dumps(source_meta,ensure_ascii=False,indent=2,sort_keys=True)+"\n")
        comp=posting_compensation.extract_file(path,tracker_id=tid,source_url=rec["url"],structured_record=rec)
        comp["posting"]["path"]=archive_rel
        comp_path=stage_meta/"compensation"/f"{tid}_PostingCompensationV1.json";posting_compensation.atomic_write(comp_path,json.dumps(comp,ensure_ascii=False,indent=2,sort_keys=True)+"\n")
        elig=posting_eligibility.extract_file(path,tracker_id=tid,source_url=rec["url"],source_id=rec["source_id"],structured_record=rec,captured_at=capture["captured_at"],capture_provenance=binding)
        elig["posting"]["path"]=archive_rel
        elig_path=stage_meta/"eligibility"/f"{tid}_PostingEligibilityV1.json";posting_eligibility.atomic_write(elig_path,json.dumps(elig,ensure_ascii=False,indent=2,sort_keys=True)+"\n")
        row=ws.max_row+1
        # Preserve the historical V1 direct-record behavior for compatibility.
        # A validated source-neutral V2 bundle no longer forces its URL into
        # the legacy LinkedIn-labelled column.
        linkedin_cell=rec["url"] if bundle_sha is None else ""
        values={"Tracker ID":tid,"Previous Row":row,"Puesto":rec["title"],"Empresa":rec["company"],"Pais":country_from_location(rec["location"]),"Estatus":"Saved","Next Action":"Triage","JD File":archive_rel,"Link puesto linkedin":linkedin_cell,"Link puesto empresa":rec["url"],"Comentarios":f"Direct-source intake {dt.date.today().isoformat()} — pending triage"}
        for key,val in values.items():ws.cell(row,hdr[key],val)
        written.append({"tracker_id":tid,"source_id":rec["source_id"],"url":rec["url"],"archive":values["JD File"],"source_sidecar":f"JobPostings/_meta/sources/{tid}_PostingSourceProvenanceV2.json","title":rec["title"],"company":rec["company"]})
    table_ref=extend_table(ws)
    fd,tmp=tempfile.mkstemp(prefix=".external-intake-",suffix=".xlsx",dir=target.parent);os.close(fd)
    try:
        wb.save(tmp);check=openpyxl.load_workbook(tmp,read_only=True,data_only=False)
        if check["Jobs"].max_row!=ws.max_row:raise ValueError("tracker row-count readback failed")
        check.close()
        if not args.dry_run and hashlib.sha256(tracker.read_bytes()).hexdigest()!=initial:raise ValueError("tracker changed concurrently")
        os.replace(tmp,target)
    finally:
        if os.path.exists(tmp):os.unlink(tmp)
    if not args.dry_run:
        publications=[]
        try:
            for source,destination in [(p,postings/p.name) for p in stage_postings.glob("*.md")]:
                if destination.exists():raise ValueError(f"commit destination already exists: {destination}")
                destination.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(source,destination);publications.append(destination)
            for source in stage_meta.rglob("*.json"):
                destination=meta/source.relative_to(stage_meta)
                if destination.exists():raise ValueError(f"commit destination already exists: {destination}")
                destination.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(source,destination);publications.append(destination)
            if hashlib.sha256(tracker.read_bytes()).hexdigest()!=initial:raise ValueError("tracker changed before commit publication")
            # ``backup_dir`` may live on another filesystem (the default is
            # /private/tmp).  Copy the verified staged workbook to a sibling
            # temporary file first so the final replace remains atomic.
            publish_fd,publish_tmp=tempfile.mkstemp(prefix=".external-intake-publish-",suffix=".xlsx",dir=tracker.parent)
            os.close(publish_fd)
            try:
                shutil.copy2(target,publish_tmp)
                if hashlib.sha256(tracker.read_bytes()).hexdigest()!=initial:raise ValueError("tracker changed during commit publication")
                os.replace(publish_tmp,tracker)
            finally:
                if os.path.exists(publish_tmp):os.unlink(publish_tmp)
        except Exception:
            if hashlib.sha256(tracker.read_bytes()).hexdigest()==initial:
                for path in reversed(publications):path.unlink(missing_ok=True)
            raise
    receipt_tracker=tracker if not args.dry_run else target
    receipt={"schema":"ExternalSourceIntakeReceiptV1","mode":"DRY_RUN" if args.dry_run else "COMMIT","created_at":dt.datetime.now(dt.timezone.utc).isoformat(),"tracker":str(receipt_tracker),"tracker_before_sha256":initial,"tracker_after_sha256":hashlib.sha256(receipt_tracker.read_bytes()).hexdigest(),"backup":str(backup),"written":written,"already_present":[r["source_id"] for r in already],"table_ref":table_ref}
    receipt_path=stage/"receipt.json";receipt_path.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps({"status":receipt["mode"],"written_count":len(written),"already_present_count":len(already),"stage":str(stage),"receipt":str(receipt_path),"rows":written},ensure_ascii=False,indent=2))
    return 0

if __name__=="__main__":raise SystemExit(main())
