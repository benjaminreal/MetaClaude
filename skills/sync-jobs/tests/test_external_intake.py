import hashlib
import json
from pathlib import Path
import sys
import unittest
from urllib.parse import urlsplit
import openpyxl

import test_workflow as fixture

SKILL=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(SKILL/"scripts"))
from acquisition_quality import validate_quality
from source_acquisition import FetchResponse,acquire_urls,canonical_json,request_id,resume_bundle,sha256_bytes,source_record_hash

class ExternalIntake(unittest.TestCase):
    setUp=fixture.Workflow.setUp
    tearDown=fixture.Workflow.tearDown
    runstage=fixture.Workflow.runstage

    def record(self):
        description="Coordinate a fully synthetic example process. This fixture ends here."
        return {"title":"Example role","company":"Example Corp","location":"Example location","url":"https://careers.example.com/jobs/42?utm_source=email","description":description,"status":"Saved","provenance":{"browser":"authorized-browser-a","method":"rendered_dom_text","captured_at":"2030-01-01T20:00:00Z","complete_text":True,"completion_evidence":"Expanded the posting and verified its final sentence.","status_certain":True},"quality_evidence":{"end_verified":True,"end_marker":"This fixture ends here.","truncation_flags":[]}}

    def fallback_bundle(self,url,provider_job_id=None,title="Synthetic fallback role"):
        shell=(SKILL/"tests/fixtures/source_acquisition/javascript_shell.html").read_bytes()
        def fetch(target):
            status=503 if target!=url else 200
            body=b"unavailable" if target!=url else shell
            return FetchResponse(target,target,status,{"content-type":"text/html"},body)
        original=acquire_urls([url],fetch)
        description="Complete synthetic browser capture with a verified ending."
        evidence_dir=self.root/"fallback-evidence";evidence_dir.mkdir(exist_ok=True)
        capture_name=hashlib.sha256(url.encode()).hexdigest()[:16]+".html"
        (evidence_dir/capture_name).write_text(f'<div id="jobDescriptionText"><p>{description}</p></div>')
        provider=urlsplit(url).hostname
        if "greenhouse.io" in provider:provider="greenhouse"
        replacements={"schema":"SourcePostingFallbackV2","outcomes":[{"request_id":request_id(url),"record":{
            "url":url,"title":title,"company":"Example Corp","location":"Remote, US","description":description,
            "provider":provider,"provider_job_id":provider_job_id,
            "provenance":{"browser":"authorized-browser-a","method":"rendered_dom_text","captured_at":"2030-01-01T00:00:00Z","complete_text":True,"completion_evidence":"Expanded and checked the final section.","status_certain":False,"status_uncertainty":"Not asserted."},
            "quality_evidence":{"end_verified":True,"end_marker":"verified ending.","truncation_flags":[],"capture_path":f"fallback-evidence/{capture_name}"},
        }}]}
        return resume_bundle(original,replacements,evidence_base=self.root)

    def batch_with_identity_collisions(self,*,same_url):
        url="https://careers.unknown.example/openings/42"
        body=(SKILL/"tests/fixtures/source_acquisition/generic_jsonld.html").read_bytes()
        bundle=acquire_urls([url],lambda target:FetchResponse(target,target,200,{"content-type":"text/html"},body))
        first=bundle["records"][0];first["provider_job_id"]="stable-42"
        first["source_identity"]={"kind":"provider_job_id","provider":first["provider"],"provider_job_id":"stable-42","canonical_url":first["canonical_url"]}
        second=json.loads(json.dumps(first));second["request_id"]="req-"+"b"*20
        if same_url:
            second["provider_job_id"]="conflicting-42"
            second["source_identity"]["provider_job_id"]="conflicting-42"
        else:
            changed=url+"?source=changed"
            second.update({"supplied_url":changed,"final_url":changed,"canonical_url":changed,"url":changed})
            second["source_identity"]["canonical_url"]=changed
        for record in (first,second):
            record["quality"]=validate_quality(record);record["source_record_sha256"]=source_record_hash(record)
        bundle["selected_requests"]=[{"request_id":first["request_id"],"supplied_url":first["supplied_url"]},{"request_id":second["request_id"],"supplied_url":second["supplied_url"]}]
        bundle["records"]=[first,second]
        bundle["bundle_sha256"]=sha256_bytes(canonical_json({key:value for key,value in bundle.items() if key!="bundle_sha256"}))
        return bundle

    def test_external_dry_run_commit_and_idempotency(self):
        payload=self.root/"external.json";payload.write_text(json.dumps({"records":[self.record()]}))
        before=fixture.digest(self.tracker)
        self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/"dry")
        self.assertEqual(fixture.digest(self.tracker),before)
        self.assertEqual(len(list((self.root/"JobPostings/postings").glob("*ext-*.md"))),0)
        self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/"live",commit=True)
        wb=openpyxl.load_workbook(self.tracker,read_only=True,data_only=True);ws=wb["Jobs"]
        self.assertEqual(ws.max_row,3);self.assertEqual(ws.cell(3,6).value,"Saved")
        self.assertEqual(ws.cell(3,9).value,self.record()["url"]);self.assertEqual(ws.cell(3,12).value,self.record()["url"]);wb.close()
        archive=next((self.root/"JobPostings/postings").glob("*ext-*.md"));self.assertIn("Source: https://careers.example.com/jobs/42?utm_source=email",archive.read_text())
        sidecar=json.loads(next((self.root/"JobPostings/_meta/eligibility").glob("J-000002_*.json")).read_text())
        self.assertTrue(sidecar["posting"]["source_id"].startswith("ext-"));self.assertIsNone(sidecar["posting"]["linkedin_id"])
        after=fixture.digest(self.tracker)
        self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/"again",commit=True)
        self.assertEqual(fixture.digest(self.tracker),after)

    def test_external_quality_and_identity_fail_closed(self):
        record=self.record();record["quality_evidence"]["end_marker"]="wrong ending"
        payload=self.root/"bad.json";payload.write_text(json.dumps({"records":[record]}));before=fixture.digest(self.tracker)
        self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/"bad",commit=True,ok=False)
        self.assertEqual(fixture.digest(self.tracker),before)

    def test_source_neutral_v2_dry_run_commit_and_sidecar(self):
        url="https://careers.unknown.example/openings/42"
        body=(SKILL/"tests/fixtures/source_acquisition/generic_jsonld.html").read_bytes()
        def fetch(target):
            return FetchResponse(target,target,200,{"content-type":"text/html; charset=utf-8"},body)
        bundle=acquire_urls([url],fetch)
        payload=self.root/"source-v2.json";payload.write_text(json.dumps(bundle))
        before=fixture.digest(self.tracker)
        self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/"v2-dry")
        self.assertEqual(fixture.digest(self.tracker),before)
        self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/"v2-live",commit=True)
        wb=openpyxl.load_workbook(self.tracker,read_only=True,data_only=True);ws=wb["Jobs"]
        self.assertIsNone(ws.cell(3,9).value)
        self.assertEqual(ws.cell(3,12).value,url)
        wb.close()
        source_path=self.root/"JobPostings/_meta/sources/J-000002_PostingSourceProvenanceV2.json"
        self.assertTrue(source_path.exists())
        source=json.loads(source_path.read_text())
        self.assertEqual(source["schema"],"PostingSourceProvenanceV2")
        self.assertEqual(source["discovery_url"],url)
        self.assertEqual(source["acquisition_bundle_sha256"],bundle["bundle_sha256"])
        eligibility=json.loads((self.root/"JobPostings/_meta/eligibility/J-000002_PostingEligibilityV1.json").read_text())
        worklist=self.root/"worklist-v2.json"
        self.runstage("worklist","--out",worklist)
        row=next(item for item in json.loads(worklist.read_text())["rows"] if item["tracker_id"]=="J-000002")
        self.assertEqual(row["job_url"],url)
        self.assertEqual(row["eligibility_evidence"]["availability"],"valid",{"worklist":row["eligibility_evidence"],"posting":eligibility.get("posting")})
        after=fixture.digest(self.tracker)
        replay=self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/"v2-replay",commit=True)
        self.assertIn('"status": "ALREADY_PRESENT"',replay.stdout);self.assertEqual(fixture.digest(self.tracker),after)

    def test_incomplete_or_tampered_v2_bundle_blocks_before_writes(self):
        url="https://www.linkedin.com/jobs/view/1234567890"
        bundle=acquire_urls([url],lambda _: self.fail("LinkedIn must not call HTTP"))
        payload=self.root/"incomplete-v2.json";payload.write_text(json.dumps(bundle));before=fixture.digest(self.tracker)
        self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/"incomplete",commit=True,ok=False)
        self.assertEqual(fixture.digest(self.tracker),before)

    def test_resumed_v2_bundle_direct_ingest_dry_run(self):
        bundle=self.fallback_bundle("https://js.example/jobs/9876543210")
        payload=self.root/"resumed.json";payload.write_text(json.dumps(bundle));before=fixture.digest(self.tracker)
        result=self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/"resumed-dry")
        self.assertIn('"written_count": 1',result.stdout);self.assertEqual(fixture.digest(self.tracker),before)

    def test_v2_stable_identity_survives_changed_url_and_rejects_conflict(self):
        first=self.fallback_bundle("https://boards.greenhouse.io/example/jobs/5550001111?source=first")
        path=self.root/"first.json";path.write_text(json.dumps(first))
        self.runstage("external-ingest","--records-json",path,"--backup-dir",self.root/"first",commit=True)
        second=self.fallback_bundle("https://boards.greenhouse.io/example/jobs/5550001111?source=changed")
        path2=self.root/"second.json";path2.write_text(json.dumps(second))
        result=self.runstage("external-ingest","--records-json",path2,"--backup-dir",self.root/"second",commit=True)
        self.assertIn('"status": "ALREADY_PRESENT"',result.stdout)
        with self.assertRaisesRegex(ValueError,"provider job ID conflicts"):
            self.fallback_bundle("https://boards.greenhouse.io/example/jobs/5550001111?source=first",provider_job_id="other-id")

    def test_v2_batch_identity_collisions_fail_before_writes(self):
        before=fixture.digest(self.tracker)
        for same_url in (True,False):
            with self.subTest(same_url=same_url):
                payload=self.root/f"collision-{same_url}.json";payload.write_text(json.dumps(self.batch_with_identity_collisions(same_url=same_url)))
                result=self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/f"collision-{same_url}",commit=True,ok=False)
                self.assertIn("canonical URL has conflicting" if same_url else "duplicate external source identity",result.stderr)
                self.assertEqual(fixture.digest(self.tracker),before)

    def test_legacy_identity_upgrade_is_explicit_for_same_and_changed_url(self):
        legacy=self.record();legacy["url"]="https://boards.greenhouse.io/example/jobs/5550001111?source=legacy"
        payload=self.root/"legacy.json";payload.write_text(json.dumps({"records":[legacy]}))
        self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/"legacy",commit=True)
        before=fixture.digest(self.tracker)
        for suffix in ("legacy","changed"):
            v2=self.fallback_bundle(f"https://boards.greenhouse.io/example/jobs/5550001111?source={suffix}")
            path=self.root/f"upgrade-{suffix}.json";path.write_text(json.dumps(v2))
            result=self.runstage("external-ingest","--records-json",path,"--backup-dir",self.root/f"upgrade-{suffix}",commit=True,ok=False)
            self.assertIn("identity_upgrade_required",result.stderr)
            self.assertNotIn("ALREADY_PRESENT",result.stdout)
            self.assertEqual(fixture.digest(self.tracker),before)

    def test_legacy_secret_urls_fail_before_staging_without_disclosure(self):
        before=fixture.digest(self.tracker);secret="s3cr3t-legacy-value"
        urls=[f"https://user:{secret}@example.com/jobs/42",f"https://example.com/jobs/42?session_token={secret}",f"https://example.com/jobs/42;jsessionid={secret}",f"https://example.com/jobs/42%3Bjsessionid={secret}"]
        for index,url in enumerate(urls):
            with self.subTest(url_kind=index):
                rec=self.record();rec["url"]=url
                payload=self.root/f"secret-{index}.json";payload.write_text(json.dumps({"records":[rec]}))
                stage=self.root/f"secret-stage-{index}"
                result=self.runstage("external-ingest","--records-json",payload,"--backup-dir",stage,commit=True,ok=False)
                self.assertNotIn(secret,result.stdout+result.stderr)
                self.assertFalse(stage.exists());self.assertEqual(fixture.digest(self.tracker),before)

    def test_v2_canonicalizer_preserves_identity_query_parameters(self):
        urls=["https://static.example/jobs/77?source=one","https://static.example/jobs/77?source=two"]
        body=(SKILL/"tests/fixtures/source_acquisition/generic_html.html").read_bytes()
        bundle=acquire_urls(urls,lambda target:FetchResponse(target,target,200,{"content-type":"text/html"},body))
        payload=self.root/"queries.json";payload.write_text(json.dumps(bundle))
        result=self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/"queries")
        self.assertIn('"written_count": 2',result.stdout)

    def test_conflicting_description_aliases_are_rejected(self):
        rec=self.record();rec["descriptionText"]="Different nonempty description."
        payload=self.root/"alias-conflict.json";payload.write_text(json.dumps({"records":[rec]}))
        self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/"alias",ok=False)

    def test_external_workbook_text_is_formula_inert(self):
        rec=self.record();rec["title"]=" \t=HYPERLINK(\"bad\")"
        payload=self.root/"formula.json";payload.write_text(json.dumps({"records":[rec]}))
        self.runstage("external-ingest","--records-json",payload,"--backup-dir",self.root/"formula",commit=True)
        wb=openpyxl.load_workbook(self.tracker,data_only=False);cell=wb["Jobs"].cell(3,3)
        self.assertEqual(cell.value," \t=HYPERLINK(\"bad\")");self.assertNotEqual(cell.data_type,"f");wb.close()

if __name__=="__main__":unittest.main()
