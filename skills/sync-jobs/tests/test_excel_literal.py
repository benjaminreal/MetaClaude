from pathlib import Path
import sys
import tempfile
import unittest

import openpyxl

SKILL=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(SKILL/"scripts"))
from excel_literal import normalize_literal_text,write_literal_text
from write_tracker import main as write_tracker_main


class ExcelLiteral(unittest.TestCase):
    def test_formula_shaped_external_text_round_trips_as_text(self):
        values=["=1+1","+cmd","-2+3","@SUM(A1:A2)"," \t=HYPERLINK(\"x\")","\x01=cmd","\u200b@cmd","Málaga","https://example.com/jobs/1"]
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"literal.xlsx";wb=openpyxl.Workbook();ws=wb.active
            for row,value in enumerate(values,1):write_literal_text(ws.cell(row,1),value)
            wb.save(path);check=openpyxl.load_workbook(path,data_only=False)
            for row,value in enumerate(values,1):
                cell=check.active.cell(row,1)
                self.assertEqual(cell.value,normalize_literal_text(value));self.assertEqual(cell.data_type,"s")
            self.assertEqual(check.active.cell(8,1).value,"Málaga")
            self.assertEqual(check.active.cell(9,1).value,"https://example.com/jobs/1")
            check.close()

    def test_system_numeric_and_date_cells_remain_typed(self):
        import datetime
        wb=openpyxl.Workbook();ws=wb.active;ws.cell(1,1).value=42;ws.cell(2,1).value=datetime.date(2030,1,2)
        self.assertEqual(ws.cell(1,1).data_type,"n");self.assertEqual(ws.cell(2,1).value,datetime.date(2030,1,2))

    def test_write_tracker_uses_literal_writer_for_untrusted_results(self):
        import json
        from openpyxl.worksheet.table import Table
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);tracker=root/"jobs.xlsx";results=root/"results.json"
            wb=openpyxl.Workbook();ws=wb.active;ws.title="Jobs";ws.append(["Tracker ID","Triage Summary"]);ws.append(["J-000001",""])
            ws.add_table(Table(displayName="JobsTable",ref="A1:B2"));wb.save(tracker)
            results.write_text(json.dumps({"report_batch":"synthetic.json","results":{"J-000001":{"cells":{"Triage Summary":"  =HYPERLINK(\"bad\")"}}}}))
            self.assertEqual(write_tracker_main(["--tracker",str(tracker),"--results",str(results),"--backup-dir",str(root)]),0)
            check=openpyxl.load_workbook(tracker,data_only=False);cell=check["Jobs"].cell(2,2)
            self.assertEqual(cell.value,"'  =HYPERLINK(\"bad\")");self.assertEqual(cell.data_type,"s");check.close()


if __name__=="__main__":unittest.main()
