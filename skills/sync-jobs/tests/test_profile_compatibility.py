"""Integration coverage for legacy and newer generic profile result schemas."""
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

import openpyxl


SOURCE_SKILL = Path(__file__).resolve().parents[1]
V5_PROFILE_ID = "synthetic-v5"
V5_PROFILE = {
    "id": V5_PROFILE_ID,
    "version": "1.0.0-synthetic",
    "policy": "triage_policy.py",
    "judgment_template": "judgment_template.json",
    "synthetic_fixture": True,
    "default_workspace": {
        "tracker_name": "jobs.xlsx",
        "sheet": "Jobs",
        "postings": "JobPostings/postings",
        "reports": "JobPostings/_meta",
    },
}


V5_POLICY = r'''import argparse
import datetime as dt
import json
import os


def derive(entry, today, batch_path):
    judgment = entry.get("judgment") or {}
    verdict = judgment.get("verdict", "Maybe")
    assess_only = bool(entry.get("assess_only")) or str(
        entry.get("estatus_existing") or ""
    ).casefold() in {"applied", "in-progress", "interview", "offer"}
    cells = {
        "Triage Date": today,
        "Triage Batch": batch_path,
        "Lane": judgment.get("lane", "Synthetic lane"),
        "Market Tier": judgment.get("market_tier", 0),
        "Eligibility": judgment.get("eligibility", "pass"),
        "Priority Score": judgment.get("priority_score", 0),
        "Match": judgment.get("match", 0),
        "Fit Score": judgment.get("raw_fit_score", 0),
        "Priority": judgment.get("priority", 0),
    }
    if not assess_only:
        cells["Estatus"] = verdict
        cells["Next Action"] = "Review synthetic result"
    return {
        "tracker_id": entry["tracker_id"],
        "verdict": verdict,
        "band": "synthetic",
        "assess_only": assess_only,
        "missing_jd": bool(entry.get("missing_jd")),
        "held": bool(judgment.get("held")),
        "reasons": [],
        "flags": ["SYNTHETIC_ASSESS_ONLY"] if assess_only else [],
        "cells": cells,
        "legacy": {"verdict": "Saved"},
    }


def main(argv=None):
    parser = argparse.ArgumentParser()
    parser.add_argument("--worklist", required=True)
    parser.add_argument("--out", required=True)
    parser.add_argument("--preview", required=True)
    parser.add_argument("--report-name")
    args = parser.parse_args(argv)
    with open(args.worklist, encoding="utf-8") as handle:
        worklist = json.load(handle)
    batch = args.report_name or "JobPostings/_meta/triage_synthetic.md"
    rows = worklist.get("rows", [])
    payload = {
        "generated": dt.datetime.now().isoformat(timespec="seconds"),
        "framework": "synthetic-v5-compatibility-test",
        "report_batch": batch,
        "source_worklist": os.path.abspath(args.worklist),
        "results": {
            row["tracker_id"]: derive(row, dt.date.today().isoformat(), batch)
            for row in rows
        },
    }
    with open(args.out, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, indent=2)
    with open(args.preview, "w", encoding="utf-8") as handle:
        handle.write("synthetic profile preview\n")
    return 0
'''


class ProfileCompatibility(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="sync-jobs-profile-compat-")
        self.root = Path(self.temp.name)
        self.skill = self.root / "skill"
        shutil.copytree(SOURCE_SKILL, self.skill,
                        ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))
        self.script = self.skill / "scripts" / "sync_jobs.py"

    def tearDown(self):
        self.temp.cleanup()

    def add_v5_profile(self):
        profile_dir = self.skill / "profiles" / V5_PROFILE_ID
        profile_dir.mkdir(parents=True)
        (profile_dir / "profile.json").write_text(json.dumps(V5_PROFILE), encoding="utf-8")
        (profile_dir / "triage_policy.py").write_text(V5_POLICY, encoding="utf-8")
        template = self.skill / "profiles" / "example" / "judgment_template.json"
        shutil.copyfile(template, profile_dir / "judgment_template.json")

    def make_tracker(self, root, headers, rows):
        root.mkdir(parents=True, exist_ok=True)
        tracker = root / "jobs.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Jobs"
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header) for header in headers])
        workbook.save(tracker)
        return tracker

    def invoke(self, root, profile, command, *args, ok=True):
        completed = subprocess.run(
            [sys.executable, "-B", str(self.script), "--project-root", str(root),
             "--profile", profile, command, *map(str, args)],
            capture_output=True, text=True, cwd=self.root,
        )
        if ok:
            self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
        else:
            self.assertNotEqual(completed.returncode, 0, completed.stdout + completed.stderr)
        return completed

    def artifacts(self, root, rows, prefix):
        worklist = root / f"{prefix}-worklist.json"
        results = root / f"{prefix}-results.json"
        preview = root / f"{prefix}-preview.txt"
        worklist.write_text(json.dumps({"selector": "synthetic test selection", "rows": rows}),
                            encoding="utf-8")
        self.invoke(root, "synthetic-v5", "score", "--worklist", worklist,
                    "--out", results, "--preview", preview)
        return worklist, results, preview

    def test_v5_fields_zero_lifecycle_discard_gate_and_scope_checks(self):
        self.add_v5_profile()
        root = self.root / "v5-workspace"
        headers = [
            "Tracker ID", "Puesto", "Empresa", "Pais", "Estatus", "Next Action",
            "JD File", "Link puesto linkedin", "Link puesto empresa", "Comentarios",
            "Track", "Fit Score", "Priority", "Sponsor Status", "Next Action Date",
            "Triage Summary", "Decision Driver", "Primary Risk / Blocker", "Key Uplifts",
            "Triage Date", "Triage Batch", "Lane", "Market Tier", "Eligibility",
            "Priority Score", "Match",
        ]
        base_rows = [
            {"Tracker ID": "J-100", "Puesto": "Synthetic zero case", "Empresa": "Example",
             "Estatus": "Saved", "Next Action": "Old action"},
            {"Tracker ID": "J-101", "Puesto": "Synthetic assess-only case", "Empresa": "Example",
             "Estatus": "Applied", "Next Action": "Preserve this"},
            {"Tracker ID": "J-102", "Puesto": "Synthetic discard case", "Empresa": "Example",
             "Estatus": "Saved", "Next Action": "Old action"},
        ]
        tracker = self.make_tracker(root, headers, base_rows)
        before = hashlib.sha256(tracker.read_bytes()).hexdigest()
        rows = [
            {"tracker_id": "J-100", "puesto": "Synthetic zero case", "empresa": "Example",
             "job_url": "https://jobs.example.test/100", "jd_file": "100.md",
             "estatus_existing": "Saved", "judgment": {"verdict": "Maybe", "lane": "Synthetic lane",
             "market_tier": 0, "eligibility": "pass", "priority_score": 0, "match": 0,
             "track": "Synthetic track", "raw_fit_score": 0, "priority": 0}},
            {"tracker_id": "J-101", "puesto": "Synthetic assess-only case", "empresa": "Example",
             "job_url": "https://jobs.example.test/101", "jd_file": "101.md",
             "estatus_existing": "Applied", "judgment": {"verdict": "Maybe", "lane": "Synthetic lane",
             "market_tier": 0, "eligibility": "held", "priority_score": 0, "match": 0,
             "raw_fit_score": 0, "priority": 0}},
            {"tracker_id": "J-102", "puesto": "Synthetic discard case", "empresa": "Example",
             "job_url": "https://jobs.example.test/102", "jd_file": "102.md",
             "estatus_existing": "Saved", "judgment": {"verdict": "Discarded", "lane": "Synthetic lane",
             "market_tier": 0, "eligibility": "fail", "priority_score": 0, "match": 0,
             "raw_fit_score": 0, "priority": 0,
             "discard_reason_plain_english": "Synthetic reason for review."}},
        ]
        worklist, results, _ = self.artifacts(root, rows, "v5")
        payload = json.loads(results.read_text(encoding="utf-8"))
        self.assertEqual(payload["profile"]["id"], V5_PROFILE_ID)
        self.assertEqual(len(payload["profile"]["policy_sha256"]), 64)
        self.assertNotEqual(Path(payload["report_batch"]).name, "triage_2026-10-08.md")

        report = root / "JobPostings" / "_meta" / "v5-report.md"
        self.invoke(root, V5_PROFILE_ID, "report", "--worklist", worklist,
                    "--results", results, "--out", report)
        report_text = report.read_text(encoding="utf-8")
        self.assertIn(payload["profile"]["policy_sha256"], report_text)
        self.assertIn("Eligibility:** pass", report_text)
        self.assertIn("Priority score:** 0", report_text)
        self.assertIn("Trial match:** 0", report_text)
        self.assertIn("Market tier:** 0", report_text)
        self.assertIn("Legacy fit score (reference):** 0", report_text)
        self.assertIn("Legacy priority (reference):** 0", report_text)
        self.assertIn("assess-only; the existing tracker lifecycle is preserved", report_text)
        self.assertIn("Held/assess-only rows retained in this report: `J-101`", report_text)
        self.assertIn("Synthetic reason for review.", report_text)
        self.assertIn("https://jobs.example.test/102", report_text)
        self.assertEqual(report.stat().st_mode & 0o777, 0o600)

        final_report = root / payload["report_batch"]
        self.invoke(root, V5_PROFILE_ID, "report", "--worklist", worklist,
                    "--results", results)
        self.assertTrue(final_report.is_file())

        original_results = results.read_text(encoding="utf-8")
        self.invoke(root, V5_PROFILE_ID, "report", "--worklist", worklist,
                    "--results", results, "--out", report, ok=False)
        self.assertEqual(report.read_text(encoding="utf-8"), report_text)

        # The selected profile binding and result scope remain enforced before a report/write.
        bad_payload = json.loads(original_results)
        bad_payload["profile"]["policy_sha256"] = "0" * 64
        results.write_text(json.dumps(bad_payload), encoding="utf-8")
        self.invoke(root, V5_PROFILE_ID, "report", "--worklist", worklist,
                    "--results", results, "--out", root / "bad-profile.md", ok=False)
        results.write_text(original_results, encoding="utf-8")
        bad_worklist = json.loads(worklist.read_text(encoding="utf-8"))
        bad_worklist["rows"].pop()
        worklist.write_text(json.dumps(bad_worklist), encoding="utf-8")
        self.invoke(root, V5_PROFILE_ID, "report", "--worklist", worklist,
                    "--results", results, "--out", root / "bad-scope.md", ok=False)
        worklist.write_text(json.dumps({"selector": "synthetic test selection", "rows": rows}),
                            encoding="utf-8")

        backup_dir = root / "dry-run-files"
        backup_dir.mkdir()
        valid_worklist = worklist.read_text(encoding="utf-8")
        changed = json.loads(valid_worklist)
        changed["rows"][0]["judgment"]["priority_score"] = 9
        worklist.write_text(json.dumps(changed), encoding="utf-8")
        failed = self.invoke(root, V5_PROFILE_ID, "write", "--results", results,
                             "--worklist", worklist, "--backup-dir", backup_dir, ok=False)
        self.assertIn("Worklist judgments differ from scored results", failed.stderr)
        worklist.write_text(valid_worklist, encoding="utf-8")

        valid_results = results.read_text(encoding="utf-8")
        changed_results = json.loads(valid_results)
        changed_results["results"]["J-100"]["cells"]["Unrecognized Field"] = "reject me"
        results.write_text(json.dumps(changed_results), encoding="utf-8")
        failed = self.invoke(root, V5_PROFILE_ID, "write", "--results", results,
                             "--worklist", worklist, "--backup-dir", backup_dir, ok=False)
        self.assertIn("Worklist judgments differ from scored results", failed.stderr)
        results.write_text(valid_results, encoding="utf-8")

        self.invoke(root, V5_PROFILE_ID, "write", "--results", results,
                    "--worklist", worklist, "--backup-dir", backup_dir)
        self.assertEqual(hashlib.sha256(tracker.read_bytes()).hexdigest(), before)
        dryrun = next(backup_dir.rglob("*.dryrun.xlsx"))
        workbook = openpyxl.load_workbook(dryrun, data_only=False)
        sheet = workbook["Jobs"]
        index = {cell.value: cell.column for cell in sheet[1]}
        row_by_id = {sheet.cell(row, index["Tracker ID"]).value: row
                     for row in range(2, sheet.max_row + 1)}
        self.assertEqual(sheet.cell(row_by_id["J-100"], index["Market Tier"]).value, 0)
        self.assertEqual(sheet.cell(row_by_id["J-100"], index["Priority Score"]).value, 0)
        self.assertEqual(sheet.cell(row_by_id["J-100"], index["Match"]).value, 0)
        self.assertEqual(sheet.cell(row_by_id["J-100"], index["Fit Score"]).value, 0)
        self.assertEqual(sheet.cell(row_by_id["J-100"], index["Priority"]).value, 0)
        self.assertEqual(sheet.cell(row_by_id["J-101"], index["Estatus"]).value, "Applied")
        self.assertEqual(sheet.cell(row_by_id["J-101"], index["Next Action"]).value, "Preserve this")
        self.assertEqual(sheet.cell(row_by_id["J-102"], index["Estatus"]).value, "Discarded")
        workbook.close()

    def test_legacy_synthetic_profile_score_write_and_report_still_work(self):
        root = self.root / "legacy-workspace"
        headers = ["Tracker ID", "Puesto", "Empresa", "Estatus", "Next Action",
                   "Track", "Fit Score", "Priority", "Triage Summary", "Triage Date", "Triage Batch"]
        tracker = self.make_tracker(root, headers, [
            {"Tracker ID": "J-200", "Puesto": "Synthetic legacy case", "Empresa": "Example",
             "Estatus": "Saved", "Next Action": "Old action"},
        ])
        before = hashlib.sha256(tracker.read_bytes()).hexdigest()
        worklist = root / "legacy-worklist.json"
        results = root / "legacy-results.json"
        preview = root / "legacy-preview.txt"
        report = root / "JobPostings" / "_meta" / "legacy-report.md"
        worklist.write_text(json.dumps({"selector": "legacy synthetic selection", "rows": [
            {"tracker_id": "J-200", "puesto": "Synthetic legacy case", "empresa": "Example",
             "estatus_existing": "Saved", "judgment": {"verdict": "Maybe", "track": "Synthetic",
             "raw_fit_score": 0, "triage_summary": "Synthetic legacy summary"}},
        ]}), encoding="utf-8")
        self.invoke(root, "example", "score", "--worklist", worklist,
                    "--out", results, "--preview", preview)
        self.invoke(root, "example", "report", "--worklist", worklist,
                    "--results", results, "--out", report)
        text = report.read_text(encoding="utf-8")
        self.assertIn("Legacy track:** Synthetic", text)
        self.assertIn("Fit score:** 0", text)
        self.assertIn("Synthetic legacy summary", text)

        backup_dir = root / "dry-run-files"
        backup_dir.mkdir()
        self.invoke(root, "example", "write", "--results", results,
                    "--worklist", worklist, "--backup-dir", backup_dir)
        self.assertEqual(hashlib.sha256(tracker.read_bytes()).hexdigest(), before)
        workbook = openpyxl.load_workbook(next(backup_dir.rglob("*.dryrun.xlsx")))
        sheet = workbook["Jobs"]
        self.assertEqual(sheet.cell(2, headers.index("Fit Score") + 1).value, 0)
        workbook.close()

    def test_ingest_routes_explicit_backup_parent_before_subcommand(self):
        root = self.root / "ingest-workspace"
        self.make_tracker(root, ["Tracker ID", "Estatus", "Puesto", "Empresa"], [])
        jobs = root / "empty-jobs.json"
        jobs.write_text("[]", encoding="utf-8")
        saved = root / "saved.json"
        saved.write_text(json.dumps({"count": 0, "total": 0, "error": None, "jobs": []}),
                         encoding="utf-8")
        backup_parent = root / "chosen-stage-parent"
        backup_parent.mkdir()

        self.invoke(root, "example", "ingest", "--jobs-json", jobs,
                    "--saved-json", saved, "--backup-dir", backup_parent)

        runs = list(backup_parent.glob("sync-ingest-*"))
        self.assertEqual(len(runs), 1)
        self.assertEqual(runs[0].stat().st_mode & 0o777, 0o700)
        self.assertTrue(any(runs[0].glob("_sync_run_*.md")))


if __name__ == "__main__":
    unittest.main()
