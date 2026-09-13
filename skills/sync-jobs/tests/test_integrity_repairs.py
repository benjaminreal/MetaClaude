import json
import os
from pathlib import Path
import sys
import unittest
import openpyxl
import test_workflow as fixture
digest=fixture.digest
SKILL=fixture.SKILL
sys.path.insert(0,str(SKILL/'scripts'))
from sync_integrity import saved_records

class Repairs(unittest.TestCase):
 setUp=fixture.Workflow.setUp
 tearDown=fixture.Workflow.tearDown
 runstage=fixture.Workflow.runstage
 def test_saved_count_mismatch_is_rejected(self):
  data=json.loads(self.saved.read_text());data.update(total=3,card='SAVED',pagination_complete=True,termination='empty_page');self.saved.write_text(json.dumps(data))
  before=digest(self.tracker)
  for cmd,extras in [('diff',['--out',self.root/'bad.json']),('reconcile',['--out-dir',self.root/'report']),('ingest',['--jobs-json',self.jobs])]:
   self.runstage(cmd,'--saved-json',self.saved,*extras,commit=cmd=='ingest',ok=False)
  self.assertEqual(digest(self.tracker),before)
  data.update(error='HTTP 500');self.saved.write_text(json.dumps(data));self.runstage('diff','--saved-json',self.saved,'--out',self.root/'error.json',ok=False)
  data.update(error=None,total=4);self.saved.write_text(json.dumps(data));self.runstage('diff','--saved-json',self.saved,'--out',self.root/'gap.json',ok=False)
 def test_archive_only_recovery_preserves_bytes(self):
  p=self.root/'JobPostings/postings/orphan.md';p.write_text('# Original archived role\nCompany: Archive Company\nLocation: Example location\nPosted: 2030-01\nSource: https://www.linkedin.com/jobs/view/200/\n\n## About the Job\nThe original full job description.\n');before=digest(p);tracker_before=digest(self.tracker)
  self.runstage('ingest','--jobs-json',self.jobs);self.assertEqual(digest(self.tracker),tracker_before);self.assertEqual(digest(p),before)
  self.runstage('ingest','--jobs-json',self.jobs,commit=True);self.assertEqual(digest(p),before)
  wb=openpyxl.load_workbook(self.tracker);self.assertEqual(wb['Jobs'].cell(3,3).value,'Original archived role');self.assertEqual(wb['Jobs'].cell(3,8).value,'JobPostings/postings/orphan.md');wb.close()
  after=digest(self.tracker);self.runstage('ingest','--jobs-json',self.jobs,commit=True);self.assertEqual(digest(self.tracker),after)
 def test_tracker_only_recovery_preserves_identity_lifecycle(self):
  self.runstage('ingest','--jobs-json',self.jobs,commit=True)
  wb=openpyxl.load_workbook(self.tracker);ws=wb['Jobs'];ws.cell(3,6).value='Applied';ws.cell(3,7).value='Watch reply';link=ws.cell(3,8).value;wb.save(self.tracker);before=[c.value for c in ws[3]];wb.close()
  (self.root/link).unlink();self.runstage('ingest','--jobs-json',self.jobs,commit=True)
  wb=openpyxl.load_workbook(self.tracker);self.assertEqual([c.value for c in wb['Jobs'][3]],before);self.assertEqual(wb['Jobs'].max_row,3);wb.close();self.assertTrue((self.root/link).is_file())
 def test_ambiguous_archive_stops_before_new_ingest(self):
  (self.root/'JobPostings/postings/duplicate.md').write_text('# duplicate\nSource: https://www.linkedin.com/jobs/view/100/\n')
  before=digest(self.tracker);self.runstage('ingest','--jobs-json',self.jobs,commit=True,ok=False);self.assertEqual(digest(self.tracker),before)
 def test_same_physical_archive_via_two_paths_is_not_ambiguous(self):
  from sync_integrity import surfaces
  original=self.root/'JobPostings/postings/existing.md';alias=self.root/'alias-existing.md'
  try:os.link(original,alias)
  except OSError:self.skipTest('hard links unavailable')
  wb=openpyxl.load_workbook(self.tracker);wb['Jobs'].cell(2,8).value='alias-existing.md';wb.save(self.tracker);wb.close()
  rows,archives=surfaces(self.tracker,self.root,self.root/'JobPostings/postings')
  self.assertEqual(rows['100']['Tracker ID'],'J-000001');self.assertTrue(os.path.samefile(archives['100'],original))
 def test_tracker_path_binds_employer_source_archive_to_linkedin_id(self):
  from sync_integrity import surfaces
  original=self.root/'JobPostings/postings/existing.md'
  text=original.read_text().replace('Source: https://www.linkedin.com/jobs/view/100/','Source: https://employer.example/jobs/example-role')
  original.write_text(text)
  wb=openpyxl.load_workbook(self.tracker);wb['Jobs'].cell(2,8).value='JobPostings/postings/existing.md';wb.save(self.tracker);wb.close()
  rows,archives=surfaces(self.tracker,self.root,self.root/'JobPostings/postings')
  self.assertEqual(rows['100']['JD File'],'JobPostings/postings/existing.md')
  self.assertEqual(archives['100'].resolve(),original.resolve())
 def test_mixed_missing_writer_batch_is_atomic(self):
  from write_tracker import main
  out=self.root/'raw_results.json';out.write_text(json.dumps({'results':{'J-000001':{'cells':{'Next Action':'Changed'}},'J-MISSING':{'cells':{'Next Action':'Changed'}}}}))
  before=digest(self.tracker)
  with self.assertRaises(SystemExit):main(['--tracker',str(self.tracker),'--results',str(out),'--backup-dir',str(self.root)])
  self.assertEqual(digest(self.tracker),before)
 def test_documented_duplicate_keeps_history(self):
  from sync_integrity import surfaces
  wb=openpyxl.load_workbook(self.tracker);ws=wb['Jobs'];values=[c.value for c in ws[2]];values[0]='J-000002';values[5]='Duplicate';values[9]='DUPLICATE of J-000001 (kept, Applied)';ws.append(values);wb.save(self.tracker);wb.close()
  before=digest(self.tracker);rows,_=surfaces(self.tracker,self.root,self.root/'JobPostings/postings');self.assertEqual(rows['100']['Tracker ID'],'J-000001');self.assertEqual(digest(self.tracker),before)
  wb=openpyxl.load_workbook(self.tracker);wb['Jobs'].cell(3,10).value='DUPLICATE of J-999999';wb.save(self.tracker);wb.close()
  with self.assertRaises(ValueError):surfaces(self.tracker,self.root,self.root/'JobPostings/postings')
if __name__=='__main__':unittest.main()
