import hashlib
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import openpyxl
from openpyxl.worksheet.table import Table

SKILL=Path(__file__).resolve().parents[1]
HEADERS=['Tracker ID','Previous Row','Puesto','Empresa','Pais','Estatus','Next Action','JD File','Link puesto linkedin','Comentarios','Track','Link puesto empresa','Fit Score','Priority','Sponsor Status','Next Action Date','Triage Summary','Decision Driver','Primary Risk / Blocker','Key Uplifts','Triage Date','Triage Batch','Aplique (Fecha)']
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
class Workflow(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory(prefix='sync-jobs-test-');self.root=Path(self.temp.name);self.tracker=self.root/'jobs.xlsx'
  wb=openpyxl.Workbook();ws=wb.active;ws.title='Jobs';ws.append(HEADERS)
  row={'Tracker ID':'J-000001','Puesto':'Existing','Empresa':'Fixture','Estatus':'Applied','Next Action':'Watch reply','Link puesto linkedin':'https://www.linkedin.com/jobs/view/100/','Comentarios':'Keep this'}
  ws.append([row.get(h) for h in HEADERS]);ws.add_table(Table(displayName='JobsTable',ref=f'A1:W2'));wb.save(self.tracker)
  archive=self.root/'JobPostings/postings';archive.mkdir(parents=True);(archive/'existing.md').write_text('# Existing\nCompany: Fixture\nSource: https://www.linkedin.com/jobs/view/100/\n\n## About the Job\nExisting job.\n')
  self.saved=self.root/'saved.json';self.saved.write_text(json.dumps({'count':2,'total':2,'error':None,'jobs':[{'id':'100','status':'Saved'},{'id':'200','status':'Saved'}]}))
  description='An operations role with documented responsibilities.'
  record={'id':'200','title':'Role','company':'Fixture','location':'Example location','url':'https://www.linkedin.com/jobs/view/200/','description':description,'status':'Saved','provenance':{'browser':'authorized-browser-a','method':'rendered_dom_text','captured_at':'2030-01-01T00:00:00Z','complete_text':True,'completion_evidence':'Expanded and read through the final section.','status_certain':True},'quality_evidence':{'end_verified':True,'end_marker':'documented responsibilities.','truncation_flags':[]}}
  self.jobs=self.root/'jobs.json';self.jobs.write_text(json.dumps([record,record]))
 def tearDown(self):self.temp.cleanup()
 def runstage(self,command,*args,commit=False,ok=True):
  cmd=[sys.executable,str(SKILL/'scripts/sync_jobs.py'),'--project-root',str(self.root),'--profile','example']
  if commit:cmd+=['--commit']
  r=subprocess.run(cmd+[command,*map(str,args)],capture_output=True,text=True)
  if ok:self.assertEqual(r.returncode,0,r.stdout+r.stderr)
  else:self.assertNotEqual(r.returncode,0)
  return r
 def test_isolation_ingest_idempotency_reconciliation(self):
  before=digest(self.tracker)
  self.runstage('ingest','--jobs-json',self.jobs,'--saved-json',self.saved)
  self.assertEqual(digest(self.tracker),before);self.assertEqual(len(list((self.root/'JobPostings/postings').glob('*.md'))),1)
  self.runstage('ingest','--jobs-json',self.jobs,'--saved-json',self.saved,commit=True)
  wb=openpyxl.load_workbook(self.tracker);self.assertEqual(wb['Jobs'].max_row,3);self.assertEqual(wb['Jobs'].cell(2,7).value,'Watch reply');wb.close()
  after=digest(self.tracker)
  self.runstage('ingest','--jobs-json',self.jobs,'--saved-json',self.saved,commit=True)
  self.assertEqual(digest(self.tracker),after)
  self.runstage('reconcile','--saved-json',self.saved,'--out-dir',self.root/'reconciliation')
  report=next((self.root/'reconciliation').glob('*.md')).read_text();self.assertIn('unsave manually (1)',report)
 def test_excel_lock_blocks_commit(self):
  (self.root/('~$'+self.tracker.name)).touch();before=digest(self.tracker)
  self.runstage('ingest','--jobs-json',self.jobs,commit=True,ok=False);self.assertEqual(digest(self.tracker),before)
 def test_triage_and_report_gate(self):
  self.runstage('ingest','--jobs-json',self.jobs,commit=True)
  work=self.root/'work.json';results=self.root/'results.json'
  self.runstage('worklist','--out',work)
  data=json.loads(work.read_text());self.assertEqual(len(data['rows']),1)
  j=data['rows'][0]['judgment'];j.update(verdict='Discarded',raw_fit_score=40,track='Example',failed_hard_gate=False,fatal_red_flag=False,triage_summary='Synthetic fixture verdict',decision_driver='Synthetic rule',primary_risk='Synthetic risk')
  work.write_text(json.dumps(data));self.runstage('score','--worklist',work,'--out',results,'--preview',self.root/'preview.txt')
  before=digest(self.tracker)
  self.runstage('write','--results',results,'--worklist',work,commit=True,ok=False);self.assertEqual(digest(self.tracker),before)
  j['discard_reason_plain_english']='This role does not offer the required scope.';work.write_text(json.dumps(data))
  self.runstage('write','--results',results,'--worklist',work);self.assertEqual(digest(self.tracker),before)
  self.runstage('write','--results',results,'--worklist',work,commit=True)
  wb=openpyxl.load_workbook(self.tracker);self.assertEqual(wb['Jobs'].cell(3,6).value,'Discarded');self.assertEqual(wb['Jobs'].cell(2,6).value,'Applied');wb.close()
 def test_synthetic_profile_requires_explicit_verdict(self):
  spec=importlib.util.spec_from_file_location('policy',SKILL/'profiles/example/triage_policy.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  row={'tracker_id':'J-1','judgment':{'verdict':'Pursue','raw_fit_score':10,'track':'Example'}}
  self.assertEqual(m.derive(row,'2030-01-01','batch')['verdict'],'Pursue')
  row['judgment']['failed_hard_gate']=True
  self.assertEqual(m.derive(row,'2030-01-01','batch')['verdict'],'Discarded')
 def test_synthetic_profile_preserves_advanced_lifecycle(self):
  spec=importlib.util.spec_from_file_location('policy',SKILL/'profiles/example/triage_policy.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  row={'tracker_id':'J-2','estatus_existing':'Applied','assess_only':True,'judgment':{'verdict':'Maybe'}}
  result=m.derive(row,'2030-01-01','batch')
  self.assertEqual(result['verdict'],'Maybe');self.assertNotIn('Estatus',result['cells']);self.assertNotIn('Next Action',result['cells'])
 def test_public_country_mapping_requires_an_explicit_code(self):
  sys.path.insert(0,str(SKILL/'scripts'));from sync_ingest import country_from_location
  self.assertIsNone(country_from_location('Example City'))
  self.assertEqual(country_from_location('Example City, ZZ'),'ZZ')
  self.assertEqual(country_from_location('country: GB'),'GB')
if __name__=='__main__':unittest.main()
